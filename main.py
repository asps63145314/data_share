import time
import  urllib.request,csv
import  pandas as pd
import  openpyxl



if __name__ == '__main__':

    url = 'https://opendata.tdcc.com.tw/getOD.ashx?id=1-5'

    path_WeekDataCSV = 'file:///'+'D:/StockData/股權分散表/WeekData/CSV DATA/20250516.csv'

    #url = path_WeekDataCSV
    webpage = urllib.request.urlopen(url)  # 開啟網頁
    data_share_CSV = csv.reader(webpage.read().decode('utf-8').splitlines())  # 讀取資料到data陣列中


    pddata = pd.Series(webpage)
    count = 0
    Stock_vl = ["", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]
    Stock_pl = ["", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]
    Total_vl = [] #量
    Total_pl = [] #人數
    Data_date = 0
    Type_Target = ""

    wb1_Stock = openpyxl.load_workbook('股權分散表_擷取當周資料.xlsx')
    s2_Stock = wb1_Stock['股權分散表_彙整']
    s3_Stock = wb1_Stock['股權分散表_人數']
    s4_Stock = wb1_Stock['股權分散表_股數']
    Xaxis_Stock = 1
    Yaxis_Stock = 2


    wb1_ETF = openpyxl.load_workbook('股權分散表_擷取當周資料.xlsx')
    s2_ETF = wb1_ETF['股權分散表_彙整']
    s3_ETF = wb1_ETF['股權分散表_人數']
    s4_ETF = wb1_ETF['股權分散表_股數']
    Xaxis_ETF = 1
    Yaxis_ETF = 2


    Stock_EX = 0
    t_start = time.time()
    for i in data_share_CSV:
        Data_date = i[0]
        if i[1] == '證券代號':
            continue
        if Stock_EX == i[1]:
            if int(i[2]) == 16: # 差異數調整，不列入
                pass

            else:
                if int(i[2]) == 17:
                    Xaxis_Stock = int(i[2])
                    Stock_vl[int(i[2])-1] = i[4]
                    Stock_pl[int(i[2])-1] = i[3]
                else:
                    Xaxis_Stock = int(i[2]) + 1
                    Stock_vl[int(i[2])] = i[4]
                    Stock_pl[int(i[2])] = i[3]




                if Type_Target == "Stock":
                    s4_Stock.cell(Yaxis_Stock, Xaxis_Stock).value = i[4]
                    s3_Stock.cell(Yaxis_Stock, Xaxis_Stock).value = i[3]

                elif Type_Target == "ETF":

                    s4_ETF.cell(Yaxis_ETF, Xaxis_Stock).value = i[4]
                    s3_ETF.cell(Yaxis_ETF, Xaxis_Stock).value = i[3]

        else: #數據跳至下一檔


            if str.isdigit(i[1].replace(" ","")) and 9 >= int(i[1][0]) >= 1 :
                Type_Target = "Stock"

            else:
                Type_Target = "ETF"

            Total_vl.append(Stock_vl)
            Stock_vl = ["", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, Type_Target]
            Stock_vl[int(i[2])] = i[4]
            Stock_vl[0] = i[1]



            Total_pl.append(Stock_pl)
            Stock_pl = ["", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, Type_Target]
            Stock_pl[int(i[2])] = i[3]
            Stock_pl[0] = i[1]



            if Type_Target == "Stock":
                Yaxis_Stock += 1
                s4_Stock.cell(Yaxis_Stock, int(i[2]) + 1).value = i[4]
                s4_Stock.cell(Yaxis_Stock, 1).value = i[1].replace(" ","")

                s3_Stock.cell(Yaxis_Stock, int(i[2]) + 1).value = i[3]
                s3_Stock.cell(Yaxis_Stock, 1).value = i[1].replace(" ","")

                print("\033[1;34m", i[0], i[1][0], s4_Stock.cell(Yaxis_Stock, 1).value, Type_Target)


            elif Type_Target == "ETF":
                Yaxis_ETF += 1
                s4_ETF.cell(Yaxis_ETF, int(i[2]) + 1).value = i[4]
                s4_ETF.cell(Yaxis_ETF, 1).value = i[1].replace(" ","")

                s3_ETF.cell(Yaxis_ETF, int(i[2]) + 1).value = i[3]
                s3_ETF.cell(Yaxis_ETF, 1).value = i[1].replace(" ","")


                print("\033[1;32m",i[0], i[1][0], s4_ETF.cell(Yaxis_ETF, 1).value , Type_Target)



        Stock_EX = i[1]  # 上一筆資料名稱
    print(time.time()-t_start)



    wb1_Stock.save('D:/StockData/股權分散表/股權分散表_' + Data_date + '.xlsx')
    wb1_ETF.save('D:/StockData/股權分散表/ETF股權分散表_' + Data_date + '.xlsx')
    ReadDoc_1 = open("./Sort_Range.txt", "r")